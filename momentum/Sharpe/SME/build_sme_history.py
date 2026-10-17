"""
build_sme_history.py
=====================
Builds History_updated_SME.xlsx (DATA + VOLUME sheets, NIFTY500 benchmark row)
for the NSE SME (Emerge) universe, in exactly the layout momentum_lib.load_prices
/ load_volume expect (same as dhan_datahq/build_history.py output).

Why a separate script: dhandata.resolve_security_id() only matches series
EQ/BE, so SME ('SM') symbols never resolve through build_history.py. This
script selects SM-series rows from the Dhan scrip master directly and calls
dh.get_daily_history_df with the security id. The shared dhandata library is
NOT modified.

Universe : NSE, SEGMENT 'E', SERIES 'SM' in api-scrip-master-detailed.csv
           (CURRENT SME listings only -> survivorship bias, see README note).
Per-symbol results are cached in SME/cache/<SYMBOL>.csv so a killed run resumes.

Usage:
    python build_sme_history.py            # 6 years back
    python build_sme_history.py 6 --limit 20   # quick trial
"""
import argparse
import datetime as dt
import sys
import time
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.cell import WriteOnlyCell

HERE = Path(__file__).resolve().parent
DHAN_DIR = Path(r"C:\Users\ameet\Documents\Github\dhan_datahq")
sys.path.insert(0, str(DHAN_DIR))
import dhandata as dh  # noqa: E402

CACHE = HERE / "cache"
OUT = HERE / "History_updated_SME.xlsx"
SCRIP_MASTER = DHAN_DIR / "api-scrip-master-detailed.csv"
MAX_CONSEC_FAIL = 15


def weekdays(start: dt.date, end: dt.date):
    d, out = start, []
    while d <= end:
        if d.weekday() < 5:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def sme_universe():
    m = pd.read_csv(SCRIP_MASTER, low_memory=False,
                    usecols=["EXCH_ID", "SEGMENT", "SECURITY_ID", "UNDERLYING_SYMBOL", "SYMBOL_NAME", "SERIES"])
    s = m[(m.EXCH_ID == "NSE") & (m.SEGMENT == "E") & (m.SERIES == "SM")]
    s = s.drop_duplicates("UNDERLYING_SYMBOL").sort_values("UNDERLYING_SYMBOL")
    return s.reset_index(drop=True)


def fetch_symbol(sec_id, from_date, to_date):
    return dh.get_daily_history_df(int(sec_id), "NSE_EQ", "EQUITY",
                                   from_date=from_date, to_date=to_date)


def main(years: int, limit):
    t0 = time.time()
    CACHE.mkdir(exist_ok=True)
    end = dt.date.today() - dt.timedelta(days=1)
    while end.weekday() >= 5:
        end -= dt.timedelta(days=1)
    start = end.replace(year=end.year - years) if not (end.month == 2 and end.day == 29) \
        else end.replace(year=end.year - years, day=28)
    dates = weekdays(start, end)
    from_date, to_date = start.isoformat(), (end + dt.timedelta(days=1)).isoformat()
    print(f"Range {from_date} -> {end} ({len(dates)} weekday cols)")

    uni = sme_universe()
    if limit:
        uni = uni.head(limit)
    print(f"SME universe: {len(uni)} NSE SM-series symbols")

    data, failed, empty = {}, {}, []
    consec_fail = 0
    for i, r in enumerate(uni.itertuples(), 1):
        sym = r.UNDERLYING_SYMBOL
        f = CACHE / f"{sym}.csv"
        if f.exists():
            df = pd.read_csv(f, index_col=0, parse_dates=True)
        else:
            try:
                df = fetch_symbol(r.SECURITY_ID, from_date, to_date)
                consec_fail = 0
            except Exception as e:  # noqa: BLE001
                failed[sym] = str(e)
                consec_fail += 1
                if consec_fail >= MAX_CONSEC_FAIL:
                    print(f"\nAborting: {MAX_CONSEC_FAIL} consecutive failures (last: {e})")
                    break
                continue
            df.to_csv(f)
        if df is None or df.empty:
            empty.append(sym)
        else:
            data[sym] = df
        if i % 25 == 0 or i == len(uni):
            print(f"  [{i}/{len(uni)}] ok={len(data)} empty={len(empty)} failed={len(failed)} "
                  f"({time.time()-t0:.0f}s)", flush=True)

    # Benchmark
    print("Fetching NIFTY 500 benchmark ...")
    nifty = dh.get_index_history_df("NIFTY 500", from_date=from_date, to_date=to_date)

    idx = pd.DatetimeIndex(pd.to_datetime(dates))

    def _row(df, col):
        s = df[col].copy()
        s.index = pd.to_datetime(s.index).normalize()
        return s.reindex(idx).fillna(0.0).round(2).tolist()

    wb = openpyxl.Workbook(write_only=True)
    wd, wv = wb.create_sheet("DATA"), wb.create_sheet("VOLUME")
    for ws in (wd, wv):
        hdr = [WriteOnlyCell(ws, value="TICKER")]
        for d in dates:
            c = WriteOnlyCell(ws, value=d)
            c.number_format = "dd-mmm-yy"
            hdr.append(c)
        ws.append(hdr)
    for sym, df in data.items():
        wd.append([sym] + _row(df, "close"))
        # RAW SHARE VOLUME: momentum_lib.compute_turnover multiplies price x
        # VOLUME-sheet value itself (matches update_stock_price_dhan.py).
        wv.append([sym] + _row(df, "volume"))
    wd.append(["NIFTY500"] + _row(nifty, "close"))
    wv.append(["NIFTY500"] + [None] * len(dates))
    wb.save(OUT)

    print(f"\nWrote {OUT}  ({len(data)} tickers)  in {time.time()-t0:.0f}s")
    if empty:
        print(f"{len(empty)} symbols returned no rows: {empty[:20]}")
    if failed:
        print(f"{len(failed)} symbols failed: {dict(list(failed.items())[:10])}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("years", type=int, nargs="?", default=6)
    ap.add_argument("--limit", type=int, default=None)
    a = ap.parse_args()
    main(a.years, a.limit)
