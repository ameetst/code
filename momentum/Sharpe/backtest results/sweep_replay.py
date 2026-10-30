"""
sweep_replay.py
===============
Replays the CURRENT live strategy (frozen copy of dashboard_config.json in the cache: NSEAll, max N,
MDTV, min price, min market cap, rank buffer) on the cache from ranking_cache_build.py, so a full
simulation takes ~a second instead of ~4 hours, then sweeps two parameters:

    hold_rank_buffer (live value from config)   x   min hold lock in calendar days (live: 28)

Rules replicated from Sharpe.py: 4 exit triggers (52H -25% immediate, no-rank/MDTV exit, REL_DD with lock,
rank > buffer with lock); entry = fill the open slots (dynamic_n - held) by walking the RANK-sorted list,
skipping held names and names failing the min-price / min-market-cap checks; inverse-vol sizing capped at
1/MAX_N (fully re-weighted every week); 0.20% friction; 6% cash yield; weekly rebalance.

    python sweep_replay.py --max-weeks 12     # replay only (use validate_replay.py to check vs the real engine)
    python sweep_replay.py                    # full grid once the cache is complete

Outputs go to 'sweep_hold_buffer_x_lock/': sweep_results.csv, sweep_pivots.xlsx (heatmaps),
surface_cagr.png / surface_maxdd.png / surface_sharpe.png.
Does not modify any live file.
"""

import os
import sys
import math
import pickle
import argparse
import datetime
import warnings

import json
import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE_DIR = os.path.join(HERE, "ranking_cache_nseall")
OUT_DIR = os.path.join(HERE, "sweep_hold_buffer_x_lock")

warnings.filterwarnings("ignore")

CFG = json.load(open(os.path.join(CACHE_DIR, "_config.json")))
START_EQUITY = 2_000_000.0
FRICTION = 0.002
LIQUID_YIELD_PA = 0.06
MAX_N = CFG["max_n"]
MAX_WT = 1.0 / MAX_N
RFR_ANNUAL = 0.07
END_DATE = datetime.date(2026, 8, 27)

LIVE_BUFFER, LIVE_LOCK, LIVE_RELDD = int(CFG["hold_rank_buffer"]), 28, CFG["rel_dd_breach_threshold"]
BUFFERS = sorted({30, 40, 50, 60, 100, 125, 150, 200, 300, 500, LIVE_BUFFER}) + [None]   # None = no rank exit
LOCKS = [0, 14, 28, 56, 84, 112, 168, 252]
NOBUF = 10 ** 9


def nan(x):
    return x is None or x != x


def load_cache(max_weeks=None):
    w = pickle.load(open(os.path.join(CACHE_DIR, "_weekly.pkl"), "rb"))
    total = len(w["valid_dates"]) - 1
    ranks = []
    for i in range(total):
        p = os.path.join(CACHE_DIR, f"rank_{i:04d}.pkl")
        if not os.path.exists(p):
            break
        ranks.append(pickle.load(open(p, "rb")))
        if max_weeks and len(ranks) >= max_weeks:
            break
    prepared = []
    for r in ranks:
        df = r["df"]
        prepared.append(dict(
            date=r["date"], dynamic_n=r["dynamic_n"], allow_new=r["allow_new"],
            order=df.index.tolist(),
            rank=df["RANK"].to_dict(), comp=df["COMPOSITE"].to_dict(), pct=df["PCT_FROM_52H"].to_dict(),
            reldd=df["REL_52H_DD"].to_dict(), resmom=df["RES_MOM"].to_dict(), vol=df["VOL"].to_dict(),
            cmp_ok=df["CMP_ELIGIBLE"].to_dict(), mcap_ok=df["MCAP_ELIGIBLE"].to_dict()))
    row = {t: k for k, t in enumerate(w["tickers"])}
    return dict(prepared=prepared, total=total, wpx=w["weekly_px"].values, row=row,
                nifty=w["nifty_w"].values, dates=w["valid_dates"])


def simulate(cache, buffer, lock, rel_dd=LIVE_RELDD, resmom_thr=None, entry_mode="fill"):
    buf = NOBUF if buffer is None else buffer
    wpx, row, nifty = cache["wpx"], cache["row"], cache["nifty"]
    cash_wk = (1.0 + LIQUID_YIELD_PA) ** (1 / 52) - 1.0
    equity, nifty_eq = START_EQUITY, START_EQUITY
    cur = {}                       # ticker -> [entry_date, weight, entry_week]
    eq, nq, held_n, turn = [], [], [], []
    spells = []
    exits = dict(h52=0, flt=0, reldd=0, rank=0)

    for i, r in enumerate(cache["prepared"]):
        t_date, dyn, allow = r["date"], r["dynamic_n"], r["allow_new"]
        rank_d, pct_d, reldd_d = r["rank"], r["pct"], r["reldd"]
        nxt = []
        for tk, st in cur.items():
            held = (t_date - st[0]).days
            rk, pc, rd = rank_d.get(tk, math.nan), pct_d.get(tk, math.nan), reldd_d.get(tk, math.nan)
            if not nan(pc) and pc < -25:
                exits["h52"] += 1
            elif nan(rk):
                exits["flt"] += 1
            elif not nan(rd) and rd < rel_dd:
                if held >= lock:
                    exits["reldd"] += 1
                else:
                    nxt.append(tk)
            elif rk > buf:
                if held >= lock:
                    exits["rank"] += 1
                else:
                    nxt.append(tk)
            else:
                nxt.append(tk)

        n_new = dyn - len(nxt)
        if entry_mode == "topn" and allow:      # pre-change rule: top dynamic_n by rank not already held, no backfill
            held_set = set(nxt)
            nxt.extend(t for t in r["order"][:dyn] if t not in held_set)
        elif allow and n_new > 0:
            held_set = set(nxt)
            cands = []
            for t in r["order"]:
                if t in held_set:
                    continue
                if not bool(r["cmp_ok"].get(t, True)) or not bool(r["mcap_ok"].get(t, True)):
                    continue
                if resmom_thr is not None:
                    rm = r["resmom"].get(t)
                    if nan(rm) or rm <= resmom_thr:
                        continue
                if len(cands) >= n_new:
                    break
                cands.append(t)
            nxt.extend(cands)

        raw = {}
        for tk in nxt:
            comp = r["comp"].get(tk, 0.0)
            v = r["vol"].get(tk, math.nan)
            raw[tk] = comp / v if not nan(v) else comp
        total_raw = sum(raw.values())
        new = {}
        for tk in nxt:
            nw = raw[tk] / total_raw if total_raw > 0 else 1.0 / len(nxt)
            capped = min(MAX_WT, nw)
            new[tk] = [cur[tk][0] if tk in cur else t_date, capped, cur[tk][2] if tk in cur else i]

        if (not allow) and not new:
            gross = cash_wk
        else:
            tot_w = sum(s[1] for s in new.values())
            cash_w = max(0.0, 1.0 - tot_w)
            if new:
                rows = [row[t] for t in new]
                w = np.array([new[t][1] for t in new])
                with np.errstate(divide="ignore", invalid="ignore"):
                    rets = wpx[rows, i + 1] / wpx[rows, i] - 1.0
                gross = float(np.nansum(rets * w)) + cash_w * cash_wk
            else:
                gross = cash_wk
            if gross != gross:
                gross = 0.0

        ab = 0.0
        for tk in set(cur) | set(new):
            ab += abs((new[tk][1] if tk in new else 0.0) - (cur[tk][1] if tk in cur else 0.0))
        net = gross - ab * FRICTION
        for tk in cur:
            if tk not in new:
                spells.append(i - cur[tk][2])

        equity *= 1 + net
        nifty_eq *= nifty[i + 1] / nifty[i]
        eq.append(equity); nq.append(nifty_eq); held_n.append(len(new)); turn.append(ab / 2.0)
        cur = new

    return dict(equity=np.array(eq), nifty=np.array(nq), held=np.array(held_n), turnover=np.array(turn),
                spells=np.array(spells), exits=exits)


def metrics(cache, sim):
    eq = np.concatenate([[START_EQUITY], sim["equity"]])
    years = (END_DATE - cache["dates"][0]).days / 365.25
    cagr = (eq[-1] / eq[0]) ** (1 / years) - 1
    mdd = float(np.min(eq / np.maximum.accumulate(eq) - 1))
    wk = eq[1:] / eq[:-1] - 1
    vol = float(np.std(wk, ddof=1) * math.sqrt(52))
    sp = sim["spells"]
    return dict(cagr=cagr, total=eq[-1] / eq[0] - 1, mdd=mdd, vol=vol,
                sharpe=(cagr - RFR_ANNUAL) / vol if vol > 0 else float("nan"), nav=eq[-1] / eq[0] * 100,
                avg_held=float(sim["held"].mean()), avg_turnover_pct=float(sim["turnover"].mean() * 100),
                n_spells=len(sp), median_hold_wks=float(np.median(sp)) if len(sp) else float("nan"),
                pct_spells_gt26wk=float((sp > 26).mean() * 100) if len(sp) else float("nan"),
                **{f"exit_{k}": v for k, v in sim["exits"].items()})


def run_grid(cache):
    rows = []
    t0 = datetime.datetime.now()
    for b in BUFFERS:
        for lk in LOCKS:
            m = metrics(cache, simulate(cache, b, lk))
            rows.append(dict(buffer=("none" if b is None else b), lock_days=lk, **m))
        print(f"  buffer {b}: done ({(datetime.datetime.now()-t0).total_seconds():.0f}s)", flush=True)
    return pd.DataFrame(rows)


def outputs(df):
    os.makedirs(OUT_DIR, exist_ok=True)
    df.to_csv(os.path.join(OUT_DIR, "sweep_results.csv"), index=False)

    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    specs = [("CAGR", "cagr", "0.0%", True), ("Max DD", "mdd", "0.0%", True), ("Sharpe", "sharpe", "0.00", True),
             ("Volatility", "vol", "0.0%", False), ("Median hold (wks)", "median_hold_wks", "0.0", True),
             ("Avg turnover pct per wk", "avg_turnover_pct", "0.0", False), ("Avg holdings", "avg_held", "0.0", True)]
    blabels = ["none" if b is None else b for b in BUFFERS]
    for title, col, fmt, hib in specs:
        ws = wb.create_sheet(title)
        ws["A1"] = "rank buffer \\ hold lock (days)"; ws["A1"].font = Font(bold=True)
        for j, lk in enumerate(LOCKS, 2):
            c = ws.cell(row=1, column=j, value=lk); c.font = Font(bold=True)
        for i, b in enumerate(blabels, 2):
            ws.cell(row=i, column=1, value=b).font = Font(bold=True)
            for j, lk in enumerate(LOCKS, 2):
                v = df[(df.buffer == b) & (df.lock_days == lk)][col].iloc[0]
                c = ws.cell(row=i, column=j, value=float(v)); c.number_format = fmt
                c.alignment = Alignment(horizontal="center")
                if b == LIVE_BUFFER and lk == LIVE_LOCK:
                    c.font = Font(bold=True, underline="single")
        lo, hi = ("F8696B", "63BE7B") if hib else ("63BE7B", "F8696B")
        rng = f"B2:{chr(65+len(LOCKS))}{1+len(blabels)}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile",
                                      mid_value=50, mid_color="FFEB84", end_type="max", end_color=hi))
        ws.column_dimensions["A"].width = 28
    wb.save(os.path.join(OUT_DIR, "sweep_pivots.xlsx"))

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    X, Y = np.meshgrid(np.arange(len(LOCKS)), np.arange(len(BUFFERS)))
    for title, col, scale, fname in [("CAGR (%)", "cagr", 100, "surface_cagr.png"),
                                     ("Max drawdown (%)", "mdd", 100, "surface_maxdd.png"),
                                     ("Sharpe ratio", "sharpe", 1, "surface_sharpe.png")]:
        Z = np.array([[df[(df.buffer == b) & (df.lock_days == lk)][col].iloc[0] * scale
                       for lk in LOCKS] for b in blabels])
        fig = plt.figure(figsize=(10, 7))
        ax = fig.add_subplot(111, projection="3d")
        ax.plot_surface(X, Y, Z, cmap="jet", edgecolor="k", linewidth=0.3, alpha=0.95)
        li, bi = LOCKS.index(LIVE_LOCK), BUFFERS.index(LIVE_BUFFER)
        pad = 0.03 * (Z.max() - Z.min())
        ax.scatter([li], [bi], [Z[bi, li] + pad], color="white", edgecolor="black", s=140, depthshade=False, zorder=10)
        ax.set_xticks(range(len(LOCKS))); ax.set_xticklabels(LOCKS)
        ax.set_yticks(range(len(BUFFERS))); ax.set_yticklabels(blabels)
        ax.set_xlabel("hold lock (days)"); ax.set_ylabel("rank buffer"); ax.set_zlabel(title)
        ax.set_title(f"{title}: rank buffer x hold lock (white dot = live: {LIVE_BUFFER} / {LIVE_LOCK}d)")
        ax.view_init(elev=28, azim=-60)
        fig.tight_layout(); fig.savefig(os.path.join(OUT_DIR, fname), dpi=150); plt.close(fig)
    print(f"Outputs written to {OUT_DIR}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--max-weeks", type=int, default=None)
    a = ap.parse_args()
    cache = load_cache(a.max_weeks)
    print(f"Cache: {len(cache['prepared'])}/{cache['total']} weeks loaded; live config: max_n={MAX_N}, "
          f"buffer={LIVE_BUFFER}, lock={LIVE_LOCK}")
    if len(cache["prepared"]) < cache["total"]:
        m = metrics(cache, simulate(cache, LIVE_BUFFER, LIVE_LOCK))
        print("Partial cache; single live-params replay only:", {k: round(v, 4) for k, v in m.items()})
        sys.exit(0)
    df = run_grid(cache)
    outputs(df)
    print(df.sort_values("cagr", ascending=False).head(10).to_string(index=False))
