"""
backtest_summary_table.py
=========================
Builds a side-by-side backtest comparison table (Excel) from one or more run
folders produced by the backtest_*.py scripts, in the layout:

    Strategy | Period | CAGR | Total Return | Max DD | Volatility | Sharpe Ratio | Final NAV

with red->green heatmap columns and "Existing" / "Proposed" row labels.

Edit RUNS below to add/remove rows, then:  python "backtest results/backtest_summary_table.py"
Metrics come from backtest_results.csv (weekly rebalance points):
  CAGR / Total Return : from Equity (NAV rebased to 100)
  Max DD              : worst peak-to-trough on the weekly equity curve (starting equity included)
  Volatility          : stdev of weekly net returns x sqrt(52)
  Sharpe Ratio        : (CAGR - RFR) / Volatility, RFR = 7% (same as Sharpe.py)
"""

import os
import datetime
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

TITLE        = "NSEAll (2,076 stocks), Dynamic 5-25 Portfolio  -  Back Test From {start}"
SECTION      = "Weekly Rebalance"
RFR_ANNUAL   = 0.07
START_EQUITY = 2_000_000.0
END_DATE     = "2026-08-27"   # last bar in History_updated_NSEAll.xlsx (final period closes here)
OUT_FILE     = os.path.join(HERE, "backtest_summary.xlsx")

# (row label shown in table, run folder, tag shown in left column: "Existing" / "Proposed" / "")
RUNS = [
    ("NSEAll LIVE (ADTV + 4 exits)",                "run_2026-09-23_21-21-38_live_baseline_nseall",  "Existing"),
    ("NSEAll LIVE + RES_MOM>1 entry gate",          "run_2026-09-23_21-23-34_live_resmom_gt1_nseall", "Proposed"),
]
BENCHMARK_LABEL = "NIFTY500 (buy & hold)"


def metrics_from_equity(equity_series, start_date, end_date):
    eq = pd.concat([pd.Series([START_EQUITY]), equity_series.reset_index(drop=True)], ignore_index=True)
    years = (end_date - start_date).days / 365.25
    total_ret = eq.iloc[-1] / eq.iloc[0] - 1.0
    cagr = (eq.iloc[-1] / eq.iloc[0]) ** (1.0 / years) - 1.0
    max_dd = (eq / eq.cummax() - 1.0).min()
    wk_ret = eq.pct_change().dropna()
    vol = wk_ret.std(ddof=1) * np.sqrt(52)
    sharpe = (cagr - RFR_ANNUAL) / vol if vol > 0 else np.nan
    nav = eq.iloc[-1] / eq.iloc[0] * 100.0
    return dict(cagr=cagr, total=total_ret, mdd=max_dd, vol=vol, sharpe=sharpe, nav=nav)


rows, start_date = [], None
bench = None
for label, folder, tag in RUNS:
    df = pd.read_csv(os.path.join(HERE, folder, "backtest_results.csv"))
    start_date = datetime.date.fromisoformat(df["Rebalance_Date"].iloc[0])
    end_date = datetime.date.fromisoformat(END_DATE)
    rows.append((label, tag, start_date, end_date, metrics_from_equity(df["Equity"], start_date, end_date)))
    if bench is None:
        bench = metrics_from_equity(df["Nifty_Equity"], start_date, end_date)
        bench_dates = (start_date, end_date)
rows.append((BENCHMARK_LABEL, "", bench_dates[0], bench_dates[1], bench))

# ── WORKBOOK ──────────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Backtest Summary"
ws.sheet_view.showGridLines = False

PEACH   = PatternFill("solid", fgColor="F8CBAD")
ORANGE  = PatternFill("solid", fgColor="FFC000")
YELLOW  = PatternFill("solid", fgColor="FFF2CC")
LBLUE   = PatternFill("solid", fgColor="BDD7EE")
GOLD_F  = Font(name="Calibri", size=11, color="7F6000")
BOLD_F  = Font(name="Calibri", size=11, bold=True, color="000000")
TAG_F   = Font(name="Calibri", size=11, bold=True, color="C00000")
thin    = Side(style="thin", color="7F7F7F")
BOX     = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER  = Alignment(horizontal="center", vertical="center")

ws.column_dimensions["A"].width = 11
for col, w in zip("BCDEFGHI", [40, 27, 11, 13, 11, 12, 13, 11]):
    ws.column_dimensions[col].width = w

ws.merge_cells("B1:I1")
ws["B1"] = TITLE.format(start=rows[0][2].strftime("%d-%m-%Y"))
ws["B1"].font, ws["B1"].fill, ws["B1"].alignment = Font(name="Calibri", size=12, bold=True), PEACH, CENTER
ws.row_dimensions[1].height = 24

headers = ["Strategy", "Period", "CAGR", "Total Return", "Max DD", "Volatility", "Sharpe Ratio", "Final NAV"]
for c, h in enumerate(headers, 2):
    cell = ws.cell(row=2, column=c, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = BOLD_F, ORANGE, CENTER, BOX

ws.merge_cells("B3:I3")
ws["B3"] = SECTION
ws["B3"].font, ws["B3"].fill, ws["B3"].alignment = BOLD_F, ORANGE, CENTER

first_data_row = 4
for r, (label, tag, sd, ed, m) in enumerate(rows, first_data_row):
    highlighted = bool(tag)
    ws.cell(row=r, column=1, value=tag or None)
    ws.cell(row=r, column=1).font, ws.cell(row=r, column=1).fill = TAG_F, LBLUE
    ws.cell(row=r, column=1).alignment = CENTER

    vals = [label, f"{sd.strftime('%d-%m-%Y')} To {ed.strftime('%d-%m-%Y')}",
            m["cagr"], m["total"], m["mdd"], m["vol"], m["sharpe"], m["nav"]]
    fmts = [None, None, "0.00%", "0.00%", "0.00%", "0.00%", "0.00", "0.00"]
    for c, (v, f) in enumerate(zip(vals, fmts), 2):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border, cell.alignment = BOX, CENTER
        if f:
            cell.number_format = f
        if c <= 3:
            cell.font = BOLD_F if highlighted and c == 2 else GOLD_F
            cell.fill = LBLUE if highlighted and c == 2 else YELLOW
last_data_row = first_data_row + len(rows) - 1

RED, YEL, GRN = "F8696B", "FFEB84", "63BE7B"
def heat(col_letter, higher_is_better=True):
    lo, hi = (RED, GRN) if higher_is_better else (GRN, RED)
    ws.conditional_formatting.add(
        f"{col_letter}{first_data_row}:{col_letter}{last_data_row}",
        ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile", mid_value=50,
                       mid_color=YEL, end_type="max", end_color=hi))
for col, hib in [("D", True), ("E", True), ("F", True), ("G", False), ("H", True), ("I", True)]:
    heat(col, hib)

# blue outer frame, like the reference layout
frame = Side(style="thick", color="0070C0")
for r in range(1, last_data_row + 1):
    for c in range(1, 10):
        cell = ws.cell(row=r, column=c)
        b = cell.border
        cell.border = Border(
            left=frame if c == 1 else b.left, right=frame if c == 9 else b.right,
            top=frame if r == 1 else b.top,  bottom=frame if r == last_data_row else b.bottom)

wb.save(OUT_FILE)
print(f"Saved -> {OUT_FILE}")
for label, tag, sd, ed, m in rows:
    print(f"{tag or '':<9} {label:<38} CAGR {m['cagr']*100:5.2f}%  Total {m['total']*100:8.2f}%  "
          f"MDD {m['mdd']*100:6.2f}%  Vol {m['vol']*100:5.2f}%  Sharpe {m['sharpe']:.2f}  NAV {m['nav']:.2f}")
